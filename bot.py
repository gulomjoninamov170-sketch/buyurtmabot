import asyncio
import logging
import io
import requests
from datetime import datetime
from aiogram import Bot, Dispatcher, types, F
from aiogram.filters import Command
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.types import ReplyKeyboardMarkup, KeyboardButton, ReplyKeyboardRemove
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as XLImage

BOT_TOKEN = "8866405159:AAHZ3g80AWy-SV_Fg6-JA8cO_Eebgb9gN9s"
ADMIN_CHAT_ID = 111079460

logging.basicConfig(level=logging.INFO)
bot = Bot(token=BOT_TOKEN)
storage = MemoryStorage()
dp = Dispatcher(storage=storage)

class OrderForm(StatesGroup):
    manager_name = State()
    client_name = State()
    product_name = State()
    date = State()
    cubage = State()
    weight = State()
    address_china = State()
    address_uzbekistan = State()
    photo = State()
    note = State()

@dp.message(Command("start"))
async def start(message: types.Message, state: FSMContext):
    await state.clear()
    await message.answer(
        "Assalomu alaykum! 👋\nBuyurtma berish uchun quyidagi tugmani bosing:",
        reply_markup=ReplyKeyboardMarkup(
            keyboard=[[KeyboardButton(text="📦 Buyurtma berish")]],
            resize_keyboard=True
        )
    )

@dp.message(F.text == "📦 Buyurtma berish")
@dp.message(Command("order"))
async def start_order(message: types.Message, state: FSMContext):
    await state.clear()
    await state.set_state(OrderForm.manager_name)
    await message.answer(
        "👨‍💼 Менежер исмини киритинг:",
        reply_markup=ReplyKeyboardRemove()
    )

@dp.message(OrderForm.manager_name)
async def get_manager(message: types.Message, state: FSMContext):
    await state.update_data(manager_name=message.text)
    await state.set_state(OrderForm.client_name)
    await message.answer("👤 Клиент исмини киритинг:")

@dp.message(OrderForm.client_name)
async def get_client(message: types.Message, state: FSMContext):
    await state.update_data(client_name=message.text)
    await state.set_state(OrderForm.product_name)
    await message.answer("1️⃣ Маҳсулот номини киритинг:\n(Наименование товара / Product Name)")

@dp.message(OrderForm.product_name)
async def get_product(message: types.Message, state: FSMContext):
    await state.update_data(product_name=message.text)
    await state.set_state(OrderForm.date)
    today = datetime.now().strftime("%d.%m.%Y")
    await message.answer(
        f"2️⃣ Сана (Дата / Date):\nБугунги сана: {today}",
        reply_markup=ReplyKeyboardMarkup(
            keyboard=[[KeyboardButton(text=today)]],
            resize_keyboard=True
        )
    )

@dp.message(OrderForm.date)
async def get_date(message: types.Message, state: FSMContext):
    await state.update_data(date=message.text)
    await state.set_state(OrderForm.cubage)
    await message.answer("3️⃣ Кубатура (куб. м):", reply_markup=ReplyKeyboardRemove())

@dp.message(OrderForm.cubage)
async def get_cubage(message: types.Message, state: FSMContext):
    await state.update_data(cubage=message.text)
    await state.set_state(OrderForm.weight)
    await message.answer("4️⃣ Оғирлик (кг / Weight):")

@dp.message(OrderForm.weight)
async def get_weight(message: types.Message, state: FSMContext):
    await state.update_data(weight=message.text)
    await state.set_state(OrderForm.address_china)
    await message.answer("5️⃣ Хитой манзили (Address China):")

@dp.message(OrderForm.address_china)
async def get_addr_china(message: types.Message, state: FSMContext):
    await state.update_data(address_china=message.text)
    await state.set_state(OrderForm.address_uzbekistan)
    await message.answer("6️⃣ Ўзбекистон манзили (Address Uzbekistan):")

@dp.message(OrderForm.address_uzbekistan)
async def get_addr_uz(message: types.Message, state: FSMContext):
    await state.update_data(address_uzbekistan=message.text)
    await state.set_state(OrderForm.photo)
    await message.answer(
        "7️⃣ Маҳсулот расмини юборинг (Фото):\n📌 Расм бўлмаса — /skip деб ёзинг"
    )

@dp.message(OrderForm.photo, F.photo)
async def get_photo(message: types.Message, state: FSMContext):
    photo_id = message.photo[-1].file_id
    await state.update_data(photo_id=photo_id)
    await state.set_state(OrderForm.note)
    await message.answer("8️⃣ Изоҳ (Примечание / Note):\n📌 Бўлмаса — /skip деб ёзинг")

@dp.message(OrderForm.photo, F.text == "/skip")
async def skip_photo(message: types.Message, state: FSMContext):
    await state.update_data(photo_id=None)
    await state.set_state(OrderForm.note)
    await message.answer("8️⃣ Изоҳ (Примечание / Note):\n📌 Бўлмаса — /skip деб ёзинг")

@dp.message(OrderForm.note)
async def get_note(message: types.Message, state: FSMContext):
    note = "" if message.text == "/skip" else message.text
    await state.update_data(note=note)
    data = await state.get_data()
    await message.answer("⏳ Буюртма тайёрланмоқда...")
    excel_file = await create_excel(data)
    await message.answer_document(
        types.BufferedInputFile(excel_file, filename="buyurtma.xlsx"),
        caption="✅ Буюртмангиз қабул қилинди!"
    )
    try:
        excel_file2 = await create_excel(data)
        await bot.send_document(
            ADMIN_CHAT_ID,
            types.BufferedInputFile(excel_file2, filename="buyurtma.xlsx"),
            caption=(
                f"📋 Янги буюртма!\n"
                f"👨‍💼 Менежер: {data.get('manager_name', '-')}\n"
                f"🧑 Клиент: {data.get('client_name', '-')}"
            )
        )
    except Exception as e:
        logging.error(f"Xato: {e}")
    await state.clear()
    await message.answer(
        "Яна буюртма бериш учун:",
        reply_markup=ReplyKeyboardMarkup(
            keyboard=[[KeyboardButton(text="📦 Buyurtma berish")]],
            resize_keyboard=True
        )
    )

async def create_excel(data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Buyurtma"
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 32
    ws.column_dimensions['C'].width = 48
    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    dark_blue = PatternFill("solid", fgColor="2F4F8F")
    mid_blue = PatternFill("solid", fgColor="4472C4")
    light_blue = PatternFill("solid", fgColor="D9E1F2")
    yellow = PatternFill("solid", fgColor="FFD700")

    def cell(row, col, value="", font=None, fill=None, align=None):
        c = ws.cell(row=row, column=col, value=value)
        if font: c.font = font
        if fill: c.fill = fill
        if align: c.alignment = align
        c.border = border
        return c

    def merge(row, c1, c2, value="", font=None, fill=None, align=None, h=25):
        ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
        c = ws.cell(row=row, column=c1, value=value)
        if font: c.font = font
        if fill: c.fill = fill
        if align: c.alignment = align
        ws.row_dimensions[row].height = h
        for col in range(c1, c2+1):
            ws.cell(row=row, column=col).border = border
        return c

    white_bold = Font(bold=True, size=12, color="FFFFFF")
    black_bold = Font(bold=True, size=12)
    small_bold = Font(bold=True, size=9)
    r = 1

    merge(r, 1, 3, "МЕНЕДЖЕР", white_bold, dark_blue, center, 25); r += 1
    cell(r, 1, "M", black_bold, light_blue, center)
    cell(r, 2, "Менеджер исми", small_bold, light_blue, left_align)
    cell(r, 3, data.get("manager_name", ""), align=left_align)
    ws.row_dimensions[r].height = 22; r += 1

    merge(r, 1, 3, "КЛИЕНТ", white_bold, mid_blue, center, 25); r += 1
    cell(r, 1, "C", black_bold, light_blue, center)
    cell(r, 2, "Клиент исми", small_bold, light_blue, left_align)
    cell(r, 3, data.get("client_name", ""), align=left_align)
    ws.row_dimensions[r].height = 22; r += 1

    merge(r, 1, 3, "订单申请表 · ORDER APPLICATION · ЗАЯВКА НА ЗАКАЗ",
          black_bold, yellow, center, 30); r += 1

    rows_data = [
        ("1", "产品名称 ·Product Name\n· Наименование товара", data.get("product_name", "")),
        ("2", "日期·Date · Дата", data.get("date", "")),
        ("3", "立方米 ·cubage ·\nКубатура", data.get("cubage", "")),
        ("4", "重量 ·Weight · Вес", data.get("weight", "")),
        ("5", "中国地址 ·Address China\n· Адрес в Китае", data.get("address_china", "")),
        ("6", "乌兹别克斯坦地址\n·Address Uzbekistan ·\nАдрес в Узбекистане", data.get("address_uzbekistan", "")),
    ]
    for num, label, value in rows_data:
        cell(r, 1, num, black_bold, light_blue, center)
        cell(r, 2, label, small_bold, light_blue, left_align)
        cell(r, 3, value, align=center)
        ws.row_dimensions[r].height = 48
        r += 1

    merge(r, 1, 3, "照片 · Photo · Фото", white_bold, dark_blue, center, 25); r += 1

    photo_start = r
    if data.get("photo_id"):
        try:
            file = await bot.get_file(data["photo_id"])
            url = f"https://api.telegram.org/file/bot{BOT_TOKEN}/{file.file_path}"
            img_bytes = requests.get(url).content
            img = XLImage(io.BytesIO(img_bytes))
            img.width = 320
            img.height = 210
            ws.merge_cells(start_row=photo_start, start_column=1,
                           end_row=photo_start + 12, end_column=3)
            ws.row_dimensions[photo_start].height = 180
            ws.add_image(img, f"A{photo_start}")
            r = photo_start + 13
        except Exception as e:
            logging.error(f"Rasm xatosi: {e}")
            ws.row_dimensions[r].height = 120; r += 1
    else:
        ws.row_dimensions[r].height = 120; r += 1

    cell(r, 1, "N", black_bold, light_blue, center)
    cell(r, 2, "备注/Note/Примечание", small_bold, light_blue, left_align)
    cell(r, 3, data.get("note", ""), align=left_align)
    ws.row_dimensions[r].height = 25

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()

async def main():
    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())